#!/usr/bin/env python3
"""Roster freshness / staleness dashboard. Run anytime to SEE how current the roster is
without trusting anyone's word. Reports as-of dates, per-row Last Verified, confidence mix,
provider-preview rows, and the known deferred research watchlist.

It does NOT discover new models on its own. Use the source-weight sheet and official
provider docs for that sweep, then regenerate from AI_MODEL_QUALITY_ROSTER_generator.py.
Usage:  python roster_freshness_check.py [--days 14] [path-to-xlsx]
"""
import sys, datetime, os
try:
    import openpyxl
except ImportError:
    sys.exit("pip install openpyxl")

DAYS = 14
args = [a for a in sys.argv[1:]]
if "--days" in args:
    i = args.index("--days"); DAYS = int(args[i+1]); del args[i:i+2]
XLSX = args[0] if args else os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                          "AI_MODEL_QUALITY_ROSTER_2026-06-07.xlsx")
TODAY = datetime.date.today()

# Known research/deferred areas to re-check before adding active rows.
DEFERRED_WATCHLIST = [
    "Mercury 2 (Inception, fastest ~790 tok/s)", "IBM Granite 4.0", "StepFun Step 3.7 Flash",
    "OpenBMB MiniCPM5", "rumored next GPT/Gemini/Grok names only after provider docs",
]

def parse_date(v):
    if isinstance(v, datetime.datetime): return v.date()
    if isinstance(v, datetime.date): return v
    try: return datetime.datetime.strptime(str(v).strip(), "%Y-%m-%d").date()
    except Exception: return None

wb = openpyxl.load_workbook(XLSX, data_only=True)
print("="*78)
print(f"ROSTER FRESHNESS CHECK   file={os.path.basename(XLSX)}   run={TODAY}   stale>{DAYS}d")
print("="*78)

ws = wb["01 LLMs"]; hdr = [c.value for c in ws[2]]
def col(n): return hdr.index(n)+1
ni, sti, lvi, cfi, bi = col("Model Name"), col("Status"), col("Last Verified"), col("Confidence"), col("Billing Class")
rows = 0; uniq = {}; conf = {}; bill = {}; stale = []; previews = []; oldest = None
for r in range(3, ws.max_row+1):
    nm = ws.cell(r, ni).value
    if not nm: continue
    rows += 1
    st = str(ws.cell(r, sti).value); cf = ws.cell(r, cfi).value
    uniq.setdefault(nm, (cf, st)); conf[cf] = conf.get(cf, 0)+1
    bl = ws.cell(r, bi).value; bill[bl] = bill.get(bl, 0)+1
    d = parse_date(ws.cell(r, lvi).value)
    if d:
        oldest = d if oldest is None else min(oldest, d)
        if (TODAY - d).days > DAYS and nm not in [x[0] for x in stale]:
            stale.append((nm, d))
    if "Preview" in st and nm not in [h[0] for h in previews]:
        previews.append((nm, st))

print(f"\nLLMs: {len(uniq)} unique / {rows} model-surface rows")
print(f"Confidence (rows): {conf}")
print(f"Billing (rows):    {bill}")
print(f"Oldest 'Last Verified' on any LLM row: {oldest}"
      + ("  <-- ALL CURRENT" if oldest and (TODAY-oldest).days <= DAYS else "  <-- STALE, run the SOP sweep"))

print(f"\n--- STALE LLMs (Last Verified > {DAYS} days old) ---")
print("  none" if not stale else "\n".join(f"  {d}  {nm}" for nm, d in sorted(stale, key=lambda x: x[1])))

print(f"\n--- PROVIDER PREVIEW rows (provider-documented; verify before high-stakes routing) ---")
print("  none" if not previews else "\n".join(f"  {st:38s} {nm}" for nm, st in previews))

print(f"\n--- DEFERRED watchlist (known current, not yet rated; confirm relevance) ---")
print("\n".join(f"  - {x}" for x in DEFERRED_WATCHLIST))

print("\nOther sheets (data rows):")
for s in wb.sheetnames:
    if s in ("00 README & Rating System", "01 LLMs"): continue
    n = wb[s].max_row - 2 if s != "08 Surfaces & Access" else wb[s].max_row - 1
    print(f"  {s}: {n}")

print("\nNEXT: if anything above says STALE, or to discover NEW models, use")
print("09 Source Catalog & Weights, verify official provider docs first, then re-run")
print("AI_MODEL_QUALITY_ROSTER_generator.py.")
print("="*78)
